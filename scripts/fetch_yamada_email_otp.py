from __future__ import annotations

import argparse
import json
import sys
import time
from datetime import datetime
from pathlib import Path

import openpyxl


ROOT_DIR = Path(__file__).resolve().parents[1]
SCRIPT_DIR = Path(__file__).resolve().parent
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

import src.config as config  # noqa: E402
from src.connections.xlsx_connection import COLUMN_ALIASES, excel_write_lock  # noqa: E402
from src.core.email_reader_imap import get_email_otp_imap  # noqa: E402
from yamada_profile_from_excel import load_row, profile_from_record, write_profile_js  # noqa: E402


YAMADA_FROM = "noreply@tpgaw.jp,noreply@ml.yamada-denki.jp"
AUTH_COLUMNS = ["auth_code", "auth_code_at"]


def normalize_header(value: object) -> str:
    text = str(value or "").strip().lower()
    return COLUMN_ALIASES.get(text, text)


def ensure_headers(ws) -> dict[str, int]:
    headers = [normalize_header(cell.value) for cell in ws[1]]
    for column in AUTH_COLUMNS:
        if column not in headers:
            ws.cell(row=1, column=len(headers) + 1, value=column)
            headers.append(column)
    return {header: index + 1 for index, header in enumerate(headers) if header}


def excel_lock_path(xlsx_path: Path) -> Path:
    return xlsx_path.parent / f".~lock.{xlsx_path.name}#"


def write_auth_code(xlsx_path: Path, sheet_name: str, row_number: int, code: str) -> dict:
    with excel_write_lock(xlsx_path):
        lock_path = excel_lock_path(xlsx_path)
        if lock_path.exists():
            raise RuntimeError(f"Excel appears to be open/locked: {lock_path}. Đóng file rồi chạy lại.")

        wb = openpyxl.load_workbook(xlsx_path)
        try:
            ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb[wb.sheetnames[0]]
            col_map = ensure_headers(ws)
            now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            ws.cell(row=row_number, column=col_map["auth_code"], value=code)
            ws.cell(row=row_number, column=col_map["auth_code_at"], value=now)

            tmp = xlsx_path.with_name(f"{xlsx_path.stem}.{row_number}.{int(time.time())}.tmp.xlsx")
            wb.save(tmp)
            tmp.replace(xlsx_path)
            return {
                "xlsx": str(xlsx_path),
                "sheet": ws.title,
                "row": row_number,
                "auth_code_at": now,
            }
        finally:
            wb.close()


def read_source(args: argparse.Namespace) -> tuple[dict, dict]:
    if not args.xlsx:
        email = args.email.strip()
        if not email:
            raise RuntimeError("Need --email or --xlsx.")
        return {"email": email}, {}

    xlsx_path = Path(args.xlsx).expanduser()
    if not xlsx_path.exists():
        raise RuntimeError(f"XLSX not found: {xlsx_path}")
    record, row_number, sheet_name = load_row(xlsx_path, args.sheet, args.row, args.email)
    profile = profile_from_record(record)
    source = {"xlsx": str(xlsx_path), "sheet": sheet_name, "row": row_number}
    return {**record, **profile}, source


def resolve_mailbox(args: argparse.Namespace, record: dict) -> tuple[str, str, str]:
    target_email = str(args.email or record.get("email") or "").strip()
    inbox = (
        args.inbox
        or str(record.get("otp_inbox") or "").strip()
        or str(getattr(config, "CATCHALL_INBOX", "") or "").strip()
        or target_email
    )
    password = (
        args.password
        or str(record.get("otp_password") or "").strip()
        or str(getattr(config, "CATCHALL_PASSWORD", "") or "").strip()
    )
    if not password and not args.no_row_password and inbox.lower() == target_email.lower():
        password = str(record.get("password") or "").strip()
    imap_host = (
        args.imap_host
        or str(record.get("otp_imap_host") or "").strip()
        or str(getattr(config, "IMAP_HOST", "") or "").strip()
    )
    return inbox, password, imap_host


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Fetch Yamada email OTP from IMAP.")
    parser.add_argument("--xlsx", default="")
    parser.add_argument("--sheet", default=str(getattr(config, "ACTIVE_SHEET", "Iclouds") or "Iclouds"))
    parser.add_argument("--row", type=int)
    parser.add_argument("--email", default="")
    parser.add_argument("--inbox", default="", help="Mailbox to login. Defaults to otp_inbox/otp_email/catchall/target email.")
    parser.add_argument("--password", default="", help="IMAP/app password. Defaults to otp_password/otp_pass/catchall_password.")
    parser.add_argument("--imap-host", default="", help="Override IMAP host, e.g. outlook.office365.com.")
    parser.add_argument("--no-row-password", action="store_true", help="Do not fallback to row password when inbox is target email.")
    parser.add_argument("--from-filter", default=YAMADA_FROM)
    parser.add_argument("--timeout", type=int, default=int(getattr(config, "EMAIL_OTP_TIMEOUT", 120)))
    parser.add_argument("--poll-interval", type=int, default=5)
    parser.add_argument("--code-pattern", default="")
    parser.add_argument("--since-ts", type=float, default=0)
    parser.add_argument("--allow-old", action="store_true", help="Allow old emails instead of only fresh messages.")
    parser.add_argument("--write-excel", action="store_true", help="Write auth_code/auth_code_at back to the Excel row.")
    parser.add_argument("--profile-js", default="", help="Also write agents/current_profile.js style file with auth_code.")
    return parser


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()

    try:
        record, source = read_source(args)
        target_email = str(args.email or record.get("email") or "").strip()
        if not target_email:
            raise RuntimeError("Missing target email.")
        inbox, password, imap_host = resolve_mailbox(args, record)
        if not inbox or not password:
            raise RuntimeError(
                "Missing IMAP inbox/password. Dùng --inbox/--password hoặc điền otp_inbox/otp_password "
                "(alias Neppi: otp_email/otp_pass)."
            )

        since_ts = args.since_ts
        if not since_ts and not args.allow_old:
            since_ts = time.time() - int(getattr(config, "EMAIL_CLOCK_SKEW_MARGIN", 180))

        code = get_email_otp_imap(
            target_email=target_email,
            otp_email=inbox,
            otp_pass=password,
            timeout=args.timeout,
            since_ts=since_ts,
            from_filter=args.from_filter,
            code_pattern=args.code_pattern,
            poll_interval=args.poll_interval,
            imap_host=imap_host,
        )
        if not code:
            raise RuntimeError("Không lấy được OTP từ email.")

        result = {
            "ok": True,
            "email": target_email,
            "inbox": inbox,
            "imap_host": imap_host,
            "from_filter": args.from_filter,
            "auth_code": code,
            "source": source,
        }

        if args.write_excel:
            if not source.get("xlsx"):
                raise RuntimeError("--write-excel cần --xlsx.")
            result["excel"] = write_auth_code(
                Path(str(source["xlsx"])),
                str(source["sheet"]),
                int(source["row"]),
                code,
            )

        if args.profile_js:
            profile = profile_from_record(record)
            profile["email"] = target_email
            profile["auth_code"] = code
            out_path = Path(args.profile_js).expanduser()
            write_profile_js(profile, out_path)
            result["profile_js"] = str(out_path)

        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0
    except Exception as exc:
        print(f"[email-otp] {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
