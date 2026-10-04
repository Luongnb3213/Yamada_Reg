from __future__ import annotations

import argparse
import json
import os
import re
import shlex
import signal
import subprocess
import sys
import threading
import time
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl


ROOT_DIR = Path(__file__).resolve().parents[1]
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from src.connections.xlsx_connection import excel_write_lock  # noqa: E402


DUMP_DOM_PYTHON = "/Users/macbook/Library/Application Support/pipx/venvs/frida-tools/bin/python"
DUMP_DOM_SCRIPT = Path("/Users/macbook/dump_dom.py")
DUMP_DOM_OUT_DIR = ROOT_DIR / "logs" / "yamada_dom"
UNRECOGNIZED_SCREEN_STATES = {"unknown", "no_webview_or_no_result", "bad_json_result"}
NO_RETRY_SCREEN_STATES = {"temporary_member_registration_in_progress"}
NO_RETRY_OTHER_DEVICE_DETAIL = "đang trong quá trình đăng ký ở máy khác"
# Bước DOM nào im lặng quá ngần này giây coi như WebView bị treo -> kill + force-kill app.
# Nhanh hơn watchdog 180s ở tầng batch; chỉ áp cho 2 bước drive WebView qua frida RPC.
DOM_IDLE_TIMEOUT_SEC = float(os.environ.get("YAMADA_DOM_IDLE_TIMEOUT_SEC", "90") or 90)


def quote_cmd(cmd: list[str]) -> str:
    return " ".join(shlex.quote(str(part)) for part in cmd)


def run_cmd(cmd: list[str], title: str) -> str:
    return run_cmd_output(cmd, title, stream=True)


def force_kill_app(device_id: str, bundle_id: str = "") -> None:
    """Force-kill app Yamada trên 1 máy sau khi WebView treo, để nó không kẹt lại trên màn hình.

    Chạy ở tiến trình riêng (frida client mới) nên không phụ thuộc session đã treo.
    """
    cmd = [sys.executable, "scripts/yamada_dom_runner.py", "--action", "kill", "--device-id", device_id]
    if bundle_id:
        cmd.extend(["--bundle-id", bundle_id])
    try:
        subprocess.run(cmd, cwd=str(ROOT_DIR), text=True, capture_output=True, timeout=30, check=False)
        print(f"[flow] Đã force-kill app Yamada trên máy {device_id} sau khi treo.", flush=True)
    except Exception as exc:
        print(f"[flow] Force-kill app trên {device_id} lỗi (bỏ qua): {exc}", flush=True)


def run_cmd_output(
    cmd: list[str],
    title: str,
    stream: bool = False,
    idle_timeout: float = 0.0,
    kill_app_device: str = "",
) -> str:
    print(f"\n--- {title} ---", flush=True)
    print("$ " + quote_cmd(cmd), flush=True)
    proc = subprocess.Popen(
        cmd,
        cwd=str(ROOT_DIR),
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        bufsize=1,
        start_new_session=True,
    )
    assert proc.stdout is not None
    lines: list[str] = []
    idle_fired = {"v": False}
    last_activity = [time.monotonic()]
    stop_watch = threading.Event()

    def _watch(p=proc, flag=idle_fired, last=last_activity, stop=stop_watch, limit=idle_timeout):
        while not stop.wait(5.0):
            if time.monotonic() - last[0] > limit:
                flag["v"] = True
                try:
                    os.killpg(os.getpgid(p.pid), signal.SIGKILL)
                except Exception:
                    try:
                        p.kill()
                    except Exception:
                        pass
                return

    watchdog = None
    if idle_timeout and idle_timeout > 0:
        watchdog = threading.Thread(target=_watch, daemon=True)
        watchdog.start()
    try:
        for line in proc.stdout:
            last_activity[0] = time.monotonic()
            lines.append(line)
            if stream:
                print(line, end="", flush=True)
        code = proc.wait()
    finally:
        stop_watch.set()
    output = "".join(lines)
    if idle_fired["v"]:
        print(
            f"[flow] {title}: im lặng > {idle_timeout:.0f}s -> coi như treo WebView, đã kill cây tiến trình.",
            flush=True,
        )
        if kill_app_device:
            force_kill_app(kill_app_device)
        raise RuntimeError(f"{title} treo (im lặng > {idle_timeout:.0f}s).")
    if code != 0:
        if not stream and output.strip():
            print(output.strip(), flush=True)
        raise RuntimeError(f"{title} lỗi exit={code}.")
    return output


def parse_json_from_output(output: str) -> Any:
    decoder = json.JSONDecoder()
    candidates: list[tuple[int, Any]] = []
    for index, char in enumerate(output):
        if char != "{":
            continue
        try:
            value, end = decoder.raw_decode(output[index:])
        except json.JSONDecodeError:
            continue
        candidates.append((end, value))
    if not candidates:
        raise RuntimeError("Không đọc được JSON output từ DOM runner.")
    return max(candidates, key=lambda item: item[0])[1]


def walk(value: Any):
    yield value
    if isinstance(value, dict):
        for child in value.values():
            yield from walk(child)
    elif isinstance(value, list):
        for child in value:
            yield from walk(child)


def needs_auth_code(result: Any) -> bool:
    for node in walk(result):
        if not isinstance(node, dict):
            continue
        if node.get("wait") == "auth_code":
            return True
        if node.get("state") == "email_auth_code_input" and node.get("action") == "need_auth_code":
            return True
    return False


def first_dom_error(result: Any) -> str:
    for node in walk(result):
        if not isinstance(node, dict):
            continue
        if node.get("ok", True) is False and node.get("wait") != "auth_code":
            reason = node.get("reason") or node.get("error") or node.get("action") or node.get("state")
            return str(reason or "DOM stopped with ok=false.")
    return ""


def no_retry_error_detail(result: Any) -> str:
    for node in walk(result):
        if not isinstance(node, dict):
            continue
        state = str(node.get("state") or "")
        action = str(node.get("action") or "")
        reason = str(node.get("reason") or "")
        if state in NO_RETRY_SCREEN_STATES or node.get("failNoRetry") is True or action == "fail_no_retry":
            return reason or NO_RETRY_OTHER_DEVICE_DETAIL
    return ""


def member_info_filled_without_submit(result: Any) -> bool:
    for node in walk(result):
        if isinstance(node, dict) and node.get("action") == "fill_member_info_no_submit":
            return True
    return False


def maybe_complete(result: Any) -> bool:
    for node in walk(result):
        if isinstance(node, dict) and node.get("state") == "maybe_complete":
            return True
    return False


def sent_auth_email(result: Any) -> bool:
    for node in walk(result):
        if isinstance(node, dict) and node.get("action") == "send_auth_email":
            return True
    return False


def registration_done(result: Any) -> bool:
    for node in walk(result):
        if not isinstance(node, dict):
            continue
        if node.get("state") == "member_register_complete":
            return True
        if node.get("action") == "launch_app_after_registration":
            return True
        if node.get("action") == "already_logged_in_home":
            return True
    return maybe_complete(result)


def already_logged_in(result: Any) -> bool:
    for node in walk(result):
        if not isinstance(node, dict):
            continue
        if node.get("state") == "app_home_logged_in":
            return True
        if node.get("action") == "already_logged_in_home":
            return True
        if node.get("already_logged_in") is True:
            return True
    return False


def dom_history(result: Any) -> list[dict]:
    if isinstance(result, dict) and isinstance(result.get("history"), list):
        return [node for node in result["history"] if isinstance(node, dict)]
    return [node for node in walk(result) if isinstance(node, dict) and "state" in node and "action" in node]


def last_dom_step(result: Any) -> dict:
    history = dom_history(result)
    return history[-1] if history else {}


def print_dom_summary(result: Any, title: str) -> None:
    history = dom_history(result)
    if not history:
        print(f"[dom] {title}: không có history.", flush=True)
        return
    print(f"[dom] {title}: {len(history)} bước", flush=True)
    for idx, step in enumerate(history, start=1):
        state = step.get("state") or "-"
        action = step.get("action") or "-"
        extra = ""
        if step.get("wait") == "auth_code":
            extra = " | cần OTP"
        if action == "fill_member_info_no_submit":
            extra = " | đã fill thông tin, dừng trước submit"
        if action == "already_logged_in_home":
            extra = " | đã đăng nhập"
        if action == "wait_timeout":
            last = step.get("last") or {}
            extra = f" | timeout ở {last.get('state') or '?'}"
        print(f"[dom] {idx}. {state} -> {action}{extra}", flush=True)


def has_unrecognized_screen(result: Any) -> bool:
    for node in walk(result):
        if not isinstance(node, dict):
            continue
        state = str(node.get("state") or "")
        action = str(node.get("action") or "")
        if state in UNRECOGNIZED_SCREEN_STATES:
            return True
        if action == "no_action" and node.get("ok", True) is False:
            return True
    return False


def stable_dump_device_id(device_id: str) -> str:
    value = str(device_id or "").strip()
    if value.lower() in ("", "auto", "all", "*"):
        return ""
    return value


def profile_email(profile: dict, row: int) -> str:
    raw = str(profile.get("email") or "").strip()
    email = raw.split("|", 1)[0].strip()
    return email or f"row_{row}"


def dump_unrecognized_screen(args: argparse.Namespace, profile: dict, reason: str) -> str:
    device_id = stable_dump_device_id(args.device_id)
    if not device_id:
        print(
            "[dump-dom] Bỏ qua dump HTML vì device-id đang là auto/all, "
            "không đảm bảo đúng máy khi chạy nhiều device.",
            flush=True,
        )
        return ""
    if not DUMP_DOM_SCRIPT.exists():
        print(f"[dump-dom] Không thấy script dump DOM: {DUMP_DOM_SCRIPT}", flush=True)
        return ""
    if not Path(DUMP_DOM_PYTHON).exists():
        print(f"[dump-dom] Không thấy Python Frida: {DUMP_DOM_PYTHON}", flush=True)
        return ""

    safe_device = re.sub(r"[^A-Za-z0-9_.-]+", "_", device_id)[:32] or "device"
    tag = f"row_{args.row}_{safe_device}_{reason}"
    cmd = [
        DUMP_DOM_PYTHON,
        str(DUMP_DOM_SCRIPT),
        "0",
        "--device-id",
        device_id,
        "--out-dir",
        str(DUMP_DOM_OUT_DIR),
        "--email",
        profile_email(profile, args.row),
        "--tag",
        tag,
    ]

    print("\n--- Dump HTML màn chưa nhận diện ---", flush=True)
    print("$ " + quote_cmd(cmd), flush=True)
    try:
        completed = subprocess.run(
            cmd,
            cwd=str(ROOT_DIR),
            text=True,
            capture_output=True,
            timeout=60,
            check=False,
        )
    except subprocess.TimeoutExpired:
        print("[dump-dom] Timeout khi dump HTML.", flush=True)
        return ""

    output = (completed.stdout or "") + (completed.stderr or "")
    if output.strip():
        print(output, end="" if output.endswith("\n") else "\n", flush=True)
    if completed.returncode != 0:
        print(f"[dump-dom] Dump HTML lỗi exit={completed.returncode}.", flush=True)
        return ""

    match = re.search(r"->\s*(.+?\.html)\s*$", output, re.MULTILINE)
    return match.group(1).strip() if match else ""


def summarize_json_result(output: str) -> Any:
    result = parse_json_from_output(output)
    return result


def ensure_status_headers(ws) -> dict[str, int]:
    headers = [str(cell.value or "").strip() for cell in ws[1]]
    if not any(headers):
        headers = ["email", "status", "error_details"]
        ws.append(headers)
    for header in ("status", "error_details"):
        if header not in headers:
            ws.cell(row=1, column=len(headers) + 1, value=header)
            headers.append(header)
    return {header: index + 1 for index, header in enumerate(headers) if header}


def write_row_status(xlsx: Path, sheet_name: str, row: int, status: str, error_details: str = "") -> None:
    with excel_write_lock(xlsx):
        lock_path = xlsx.parent / f".~lock.{xlsx.name}#"
        if lock_path.exists():
            print(f"[excel] Không ghi được status vì Excel đang mở/lock: {lock_path}", flush=True)
            return
        wb = openpyxl.load_workbook(xlsx)
        try:
            if sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
            elif "Iclouds" in wb.sheetnames:
                ws = wb["Iclouds"]
            elif "Gmails" in wb.sheetnames:
                ws = wb["Gmails"]
            elif "Outlooks" in wb.sheetnames:
                ws = wb["Outlooks"]
            else:
                ws = wb[wb.sheetnames[0]]
            col = ensure_status_headers(ws)
            ws.cell(row=row, column=col["status"], value=status)
            ws.cell(row=row, column=col["error_details"], value=error_details)
            tmp = xlsx.with_name(f"{xlsx.stem}.{os.getpid()}.tmp.xlsx")
            wb.save(tmp)
            tmp.replace(xlsx)
        finally:
            wb.close()


def kill_app_best_effort(args: argparse.Namespace) -> None:
    """Cuối mỗi row: kill app Yamada để xả webview (~150MB) ngay, cho compressor
    hạ nhiệt trong khoảng nghỉ trước row sau. Không chờ tới row sau mới kill (như
    kill-then-spawn cũ) để tránh máy luôn ôm 1 webview nặng -> cạn RAM -> launchd
    panic -> reboot mất jailbreak. Best-effort: lỗi thì bỏ qua, không làm hỏng row."""
    cmd = [
        sys.executable,
        "scripts/yamada_dom_runner.py",
        "--action",
        "kill",
        "--device-id",
        args.device_id,
        "--ensure-killed",
    ]
    for attempt in (1, 2):
        try:
            done = subprocess.run(cmd, cwd=str(ROOT_DIR), text=True, capture_output=True, timeout=40, check=False)
        except Exception as exc:
            print(f"[flow] Kill app cuối row lỗi tạm (bỏ qua): {exc}", file=sys.stderr, flush=True)
            return
        try:
            stopped = bool(parse_json_from_output(done.stdout or "").get("stopped"))
        except Exception:
            stopped = done.returncode == 0
        if stopped:
            print("[flow] Đã kill app Yamada cuối row để xả RAM.", flush=True)
            return
        if attempt == 1:
            print("[flow] App Yamada chưa tắt sau kill, thử lại lần 2...", flush=True)
    print(
        "[flow] CẢNH BÁO: app Yamada có thể vẫn còn trên màn sau 2 lần kill cuối row.",
        file=sys.stderr,
        flush=True,
    )


def build_dom_cmd(args: argparse.Namespace, *, fresh_launch: bool = False) -> list[str]:
    cmd = [
        sys.executable,
        "scripts/yamada_dom_runner.py",
        "--action",
        "run",
        "--wait-timeout-ms",
        str(args.wait_timeout_ms),
        "--max-steps",
        str(args.max_steps),
        "--device-id",
        args.device_id,
        "--profile-js",
        str(args.profile_js),
    ]
    if args.no_submit:
        cmd.append("--no-submit")
    if fresh_launch:
        cmd.append("--fresh-launch")
    return cmd


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Run one Yamada row end-to-end, including email OTP when needed.")
    parser.add_argument("--xlsx", required=True)
    parser.add_argument("--sheet", default="Iclouds")
    parser.add_argument("--row", type=int, required=True)
    parser.add_argument("--no-reload", action="store_true", help="Deprecated/default. Do not ask Crane to reload the app after switching container.")
    parser.add_argument("--reload-crane", action="store_true", help="Opt in to Crane reloadApplicationWithIdentifier after switching container.")
    parser.add_argument("--no-submit", action="store_true")
    parser.add_argument("--device-id", default=os.environ.get("FRIDA_DEVICE_ID", "auto"))
    parser.add_argument(
        "--container-mode",
        choices=["active-then-create", "active-then-next", "next-active", "create"],
        default="create",
    )
    parser.add_argument("--profile-js", default="")
    parser.add_argument("--wait-timeout-ms", type=int, default=15000)
    parser.add_argument("--max-steps", type=int, default=20)
    parser.add_argument(
        "--dom-idle-timeout-sec",
        type=float,
        default=DOM_IDLE_TIMEOUT_SEC,
        help="Bước DOM im lặng quá ngần này giây coi như treo WebView: kill tiến trình + force-kill app. 0 = tắt.",
    )
    return parser


def main() -> int:
    args = build_parser().parse_args()
    xlsx = Path(args.xlsx).expanduser()
    row_args = ["--xlsx", str(xlsx), "--sheet", args.sheet, "--row", str(args.row)]
    profile: dict[str, Any] = {}
    last_dom_result: Any = None
    if args.profile_js:
        profile_js = Path(args.profile_js).expanduser()
    else:
        safe_device = re.sub(r"[^A-Za-z0-9_.-]+", "_", args.device_id or "auto")[:48]
        profile_js = ROOT_DIR / "agents" / "runtime" / f"current_profile_{safe_device}_r{args.row}_{os.getpid()}.js"
    args.profile_js = str(profile_js)

    try:
        crane_cmd = [
            sys.executable,
            "scripts/crane_container_manager.py",
            "ensure-row",
            *row_args,
            "--device-id",
            args.device_id,
            "--container-mode",
            args.container_mode,
        ]
        if not args.reload_crane:
            crane_cmd.append("--no-reload")
        crane_output = run_cmd_output(crane_cmd, "Chuẩn bị container")
        crane_result = summarize_json_result(crane_output)
        print(
            "[crane] container="
            f"{crane_result.get('crane_container_name') or crane_result.get('active_container_name') or crane_result.get('crane_container_label') or crane_result.get('active_container_label') or ''} "
            f"({crane_result.get('crane_container_id') or crane_result.get('active_container_id') or ''})",
            flush=True,
        )

        profile_output = run_cmd_output(
            [
                sys.executable,
                "scripts/yamada_profile_from_excel.py",
                *row_args,
                "--out",
                str(profile_js),
            ],
            "Đọc data từ Excel",
        )
        profile_result = summarize_json_result(profile_output)
        profile_value = profile_result.get("profile") if isinstance(profile_result, dict) else {}
        profile = profile_value if isinstance(profile_value, dict) else {}
        print(f"[excel] email={profile.get('email') or ''} row={args.row}", flush=True)

        otp_request_since_ts = float(int(time.time()))
        first_output = run_cmd_output(
            build_dom_cmd(args, fresh_launch=True),
            "Chạy DOM",
            idle_timeout=args.dom_idle_timeout_sec,
            kill_app_device=args.device_id,
        )
        first_result = parse_json_from_output(first_output)
        last_dom_result = first_result
        print_dom_summary(first_result, "lượt đầu")
        error = first_dom_error(first_result)
        if error:
            raise RuntimeError(error)

        final_result = first_result
        if needs_auth_code(first_result):
            print("\n[flow] App đang chờ OTP. Tự lấy OTP từ mailbox trong Excel...", flush=True)
            otp_cmd = [
                sys.executable,
                "scripts/fetch_yamada_email_otp.py",
                *row_args,
            ]
            if sent_auth_email(first_result):
                otp_cmd.extend(["--since-ts", str(otp_request_since_ts)])
            otp_cmd.extend([
                "--profile-js",
                str(profile_js),
            ])
            otp_output = run_cmd_output(
                otp_cmd,
                "Lấy OTP email",
            )
            otp_result = summarize_json_result(otp_output)
            print(f"[email] OTP={otp_result.get('auth_code') or '(không có)'}", flush=True)
            second_output = run_cmd_output(
                build_dom_cmd(args),
                "Chạy tiếp sau OTP",
                idle_timeout=args.dom_idle_timeout_sec,
                kill_app_device=args.device_id,
            )
            second_result = parse_json_from_output(second_output)
            last_dom_result = second_result
            print_dom_summary(second_result, "sau OTP")
            final_result = second_result
            if needs_auth_code(second_result):
                raise RuntimeError("Đã lấy OTP nhưng app vẫn đang chờ auth_code.")
            error = first_dom_error(second_result)
            if error:
                raise RuntimeError(error)

        if not registration_done(final_result):
            last = last_dom_step(final_result)
            raise RuntimeError(
                "DOM dừng trước màn thông tin/OTP: "
                f"state={last.get('state') or '?'} action={last.get('action') or '?'}"
            )

        status = "LOGG_IN" if already_logged_in(final_result) else "SUCCESS"
        write_row_status(xlsx, args.sheet, args.row, status, "")
        print(f"\n[flow] Xong lượt chạy row {args.row} lúc {datetime.now():%Y-%m-%d %H:%M:%S}", flush=True)
        return 0
    except Exception as exc:
        message = str(exc)
        if "Lấy OTP email lỗi" in message:
            message = "Không lấy được OTP email. Kiểm tra otp_email/otp_pass hoặc mail OTP chưa về."
        status = "FAILED"
        no_retry_detail = no_retry_error_detail(last_dom_result)
        if no_retry_detail:
            status = "FAIL_NO_RETRY"
            message = no_retry_detail
        if has_unrecognized_screen(last_dom_result):
            dump_path = dump_unrecognized_screen(args, profile, "unknown_screen")
            if dump_path:
                message = f"{message} | dumped_html={dump_path}"
        write_row_status(xlsx, args.sheet, args.row, status, message)
        print(f"[flow] Lỗi: {message}", file=sys.stderr, flush=True)
        return 1
    finally:
        kill_app_best_effort(args)


if __name__ == "__main__":
    raise SystemExit(main())
