from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl


ROOT_DIR = Path(__file__).resolve().parents[1]
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from src.connections.xlsx_connection import COLUMN_ALIASES, normalize_status  # noqa: E402


PROFILE_KEYS = [
    "email",
    "pin",
    "phone",
    "last_name",
    "first_name",
    "last_name_kana",
    "first_name_kana",
    "postal_code",
    "prefecture",
    "city",
    "address_rest",
    "dob",
    "gender",
    "password",
    "auth_code",
    "otp_inbox",
    "otp_password",
    "otp_imap_host",
    "crane_container_id",
    "crane_container_name",
    "crane_status",
]

EXTRA_ALIASES = {
    "auth code": "auth_code",
    "auth_code": "auth_code",
    "email code": "auth_code",
    "email_code": "auth_code",
    "otp": "auth_code",
    "inputcode": "auth_code",
    "認証コード": "auth_code",
    "katakana_last_name": "last_name_kana",
    "katakana first name": "first_name_kana",
    "katakana_first_name": "first_name_kana",
    "katakana last name": "last_name_kana",
    "birth date": "dob",
    "birth_date": "dob",
    "address": "address_rest",
}


def normalize_header(value: object) -> str:
    text = str(value or "").strip().lower()
    return EXTRA_ALIASES.get(text) or COLUMN_ALIASES.get(text, text)


def cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_digits(value: str) -> str:
    return re.sub(r"\D", "", value or "")


def normalize_dob(value: str) -> str:
    raw = str(value or "").strip()
    if not raw:
        return ""
    compact = normalize_digits(raw)
    if len(compact) == 8:
        return compact

    parts = re.match(r"^\s*(\d{1,4})[\/\-.](\d{1,2})[\/\-.](\d{1,4})\s*$", raw)
    if not parts:
        return compact

    first, second, third = parts.groups()
    if len(first) == 4:
        year, month, day = first, second, third
    elif len(third) == 4:
        # Spreadsheet exports like 1/1/1990 are common in the user's data.
        month, day, year = first, second, third
    else:
        return compact
    return f"{int(year):04d}{int(month):02d}{int(day):02d}"


def read_config_xlsx() -> str:
    config_path = ROOT_DIR / "config.json"
    if not config_path.exists():
        return ""
    try:
        data = json.loads(config_path.read_text(encoding="utf-8"))
    except Exception:
        return ""
    return str(data.get("xlsx_path") or "").strip()


def read_config_sheet() -> str:
    config_path = ROOT_DIR / "config.json"
    if not config_path.exists():
        return "Iclouds"
    try:
        data = json.loads(config_path.read_text(encoding="utf-8"))
    except Exception:
        return "Iclouds"
    return str(data.get("active_sheet") or "Iclouds").strip() or "Iclouds"


def load_row(xlsx_path: Path, sheet_name: str, row_number: int | None, email: str = "") -> tuple[dict, int, str]:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    try:
        if sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
        elif "Iclouds" in wb.sheetnames:
            ws = wb["Iclouds"]
        elif "Gmails" in wb.sheetnames:
            ws = wb["Gmails"]
        elif "Outlooks" in wb.sheetnames:
            ws = wb["Outlooks"]
        elif "Inputs" in wb.sheetnames:
            ws = wb["Inputs"]
        else:
            ws = wb[wb.sheetnames[0]]

        rows = ws.iter_rows(values_only=True)
        try:
            header_row = next(rows)
        except StopIteration:
            raise RuntimeError(f"Sheet {ws.title!r} is empty.")

        headers = [normalize_header(cell) for cell in header_row]
        wanted_email = email.strip().lower()

        for index, row in enumerate(rows, start=2):
            if row_number is not None and index != row_number:
                continue
            record = {}
            for pos, header in enumerate(headers):
                if not header:
                    continue
                value = cell_text(row[pos]) if pos < len(row) else ""
                if header in record and not value:
                    continue
                record[header] = value
            row_email = str(record.get("email") or "").strip()
            if not row_email:
                continue
            if wanted_email and row_email.lower() != wanted_email:
                continue
            if row_number is None and not wanted_email:
                status = normalize_status(record.get("status", "PENDING"))
                if status not in ("PENDING", "FAILED"):
                    continue
            return record, index, ws.title
    finally:
        wb.close()

    if row_number is not None:
        raise RuntimeError(f"No usable data found at row {row_number}.")
    if email:
        raise RuntimeError(f"No usable row found for email {email!r}.")
    raise RuntimeError("No PENDING/FAILED row with email was found.")


def profile_from_record(record: dict) -> dict:
    profile = {}
    for key in PROFILE_KEYS:
        value = str(record.get(key) or "").strip()
        if value:
            profile[key] = value

    if "dob" in profile:
        profile["dob"] = normalize_dob(profile["dob"])
    if "phone" in profile:
        profile["phone"] = normalize_digits(profile["phone"])
    if "postal_code" in profile:
        profile["postal_code"] = normalize_digits(profile["postal_code"])
    if "pin" in profile:
        profile["pin"] = normalize_digits(profile["pin"])
    if "gender" in profile:
        lower = profile["gender"].strip().lower()
        if lower in ("m", "male", "man", "nam", "1", "男性"):
            profile["gender"] = "male"
        elif lower in ("f", "female", "woman", "nu", "nữ", "2", "女性"):
            profile["gender"] = "female"
    return profile


def write_profile_js(profile: dict, out_path: Path) -> None:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    payload = json.dumps(profile, ensure_ascii=False, indent=2)
    js = (
        "// Generated by scripts/yamada_profile_from_excel.py\n"
        "(function loadYamadaExcelProfile(){\n"
        f"  const profile = {payload};\n"
        "  globalThis.__YAMADA_PROFILE__ = profile;\n"
        "  function apply(){\n"
        "    if (typeof setYamadaProfile === 'function') {\n"
        "      setYamadaProfile(profile);\n"
        "      console.log('[yamada-profile] loaded from Excel: ' + (profile.email || '(no email)'));\n"
        "      return;\n"
        "    }\n"
        "    setTimeout(apply, 250);\n"
        "  }\n"
        "  apply();\n"
        "})();\n"
    )
    out_path.write_text(js, encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description="Export one Excel row into a Frida-loaded Yamada profile JS file.")
    parser.add_argument("--xlsx", default=read_config_xlsx(), help="Path to XLSX. Defaults to config.json xlsx_path.")
    parser.add_argument("--sheet", default=read_config_sheet(), help="Sheet name. Defaults to config active_sheet.")
    parser.add_argument("--row", type=int, default=None, help="Excel row number to export, e.g. 2.")
    parser.add_argument("--email", default="", help="Find row by email instead of first pending row.")
    parser.add_argument("--out", default=str(ROOT_DIR / "agents" / "current_profile.js"), help="Output JS file.")
    parser.add_argument("--print-only", action="store_true", help="Print profile JSON without writing JS.")
    args = parser.parse_args()

    if not args.xlsx:
        raise SystemExit("Missing --xlsx and config.json does not contain xlsx_path.")
    xlsx_path = Path(args.xlsx).expanduser()
    if not xlsx_path.exists():
        raise SystemExit(f"XLSX not found: {xlsx_path}")

    record, row_number, sheet = load_row(xlsx_path, args.sheet, args.row, args.email)
    profile = profile_from_record(record)
    # Never trust a previously saved Excel OTP at the start of a flow. The
    # fresh code must be fetched from email after Yamada sends it for this run.
    profile.pop("auth_code", None)
    metadata = {"xlsx": str(xlsx_path), "sheet": sheet, "row": row_number}

    if args.print_only:
        print(json.dumps({"profile": profile, "source": metadata}, ensure_ascii=False, indent=2))
        return 0

    out_path = Path(args.out).expanduser()
    write_profile_js(profile, out_path)
    print(f"Wrote: {out_path}")
    print(json.dumps({"profile": profile, "source": metadata}, ensure_ascii=False, indent=2))
    print()
    print("Run Frida with:")
    print(
        '"/Users/macbook/Library/Application Support/pipx/venvs/frida-tools/bin/frida" '
        f'-U -n yamadadenki '
        f'-l {ROOT_DIR / "agents" / "yamada_register_agent.js"} '
        f'-l {out_path}'
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
